#!/usr/bin/env python3
"""Read and validate one fixed snapshot of local FarmBoy metric records."""

from __future__ import annotations

import json
import os
import re
import stat
import sys
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Any

REPOSITORY_ROOT = Path(__file__).resolve().parents[3]
CONTROLLER_SCRIPTS = REPOSITORY_ROOT / "controllers/farmboy/scripts/python"
_LAYOUT_ERROR: str | None = None
try:
    if not CONTROLLER_SCRIPTS.is_dir():
        raise ImportError("controller metrics helper is missing")
    sys.path.insert(0, str(CONTROLLER_SCRIPTS))
    from metrics import MAX_INPUT_BYTES, validate_record
except ImportError:
    MAX_INPUT_BYTES = 65_536
    validate_record = None
    _LAYOUT_ERROR = "report source layout is incomplete: controller metrics helper is unavailable"


CANDIDATE_NAME = re.compile(r"metric-([0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12})\.json\Z")
INVALID_CATEGORIES = ("unreadable", "oversized", "malformed-json", "unsupported-schema", "invalid-contract")
PROJECTION_FIELDS = (
    "schema_version", "invocation_id", "run_id", "step_id", "invocation_kind",
    "agent_name", "skill_id", "workflow_id", "project_label", "started_at",
    "finished_at", "duration_ms", "model", "effort",
)
WORKBOOK_FIELDS = (
    "invocation_id", "run_id", "project_label", "skill_id", "workflow_id", "step_id",
    "invocation_kind", "agent_name", "model", "effort", "started_at", "finished_at", "duration_ms",
)
WORKBOOK_HEADERS = (
    "invocation_id", "run_id", "project", "skill", "workflow", "step", "invocation_kind",
    "agent", "model", "effort", "started_at", "finished_at", "duration_ms",
)
MAX_ACCEPTED_ROWS = 1_000_000


class ReportInputError(RuntimeError):
    """A report input directory could not be enumerated safely."""


class _MalformedMetricJson(ValueError):
    """A candidate has invalid JSON structure or prohibited JSON values."""


class _UnparsedJsonInteger:
    """A syntactically valid integer too large for Python's guarded conversion."""


def _require_source_layout() -> None:
    """Fail with a concise error when the maintainer source tree is incomplete."""
    if _LAYOUT_ERROR is not None:
        raise RuntimeError(_LAYOUT_ERROR)


def _require_report_dependencies() -> None:
    """Check approved workbook dependencies before creating report outputs."""
    try:
        import et_xmlfile  # noqa: F401
        import openpyxl  # noqa: F401
    except ImportError as error:
        raise RuntimeError("approved report dependencies are unavailable: openpyxl==3.1.5 and et-xmlfile are required") from error
    if openpyxl.__version__ != "3.1.5":
        raise RuntimeError("approved report dependencies are unavailable: openpyxl==3.1.5 and et-xmlfile are required")


def _parse_json_integer(value: str) -> int | _UnparsedJsonInteger:
    """Parse ordinary integers while keeping oversized unknown values ignorable."""
    try:
        return int(value)
    except ValueError:
        return _UnparsedJsonInteger()


def _unique_json_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    """Build one JSON object while rejecting duplicate keys."""
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise _MalformedMetricJson("duplicate-key")
        result[key] = value
    return result


def _reject_nonfinite(value: str) -> None:
    """Reject JSON's non-standard NaN and infinity constants."""
    raise _MalformedMetricJson("nonfinite-number")


def _completed_name(name: str) -> bool:
    """Return whether a direct entry uses the completed metric filename convention."""
    return CANDIDATE_NAME.fullmatch(name) is not None


def discover_candidates(directory: Path) -> list[Path]:
    """Capture one sorted top-level list of completed regular metric files."""
    root = Path(directory)
    try:
        root_stat = root.lstat()
    except FileNotFoundError:
        return []
    except OSError as error:
        raise ReportInputError("input-inaccessible") from error
    if stat.S_ISLNK(root_stat.st_mode) or not stat.S_ISDIR(root_stat.st_mode):
        raise ReportInputError("input-inaccessible")
    try:
        entries = list(root.iterdir())
    except OSError as error:
        raise ReportInputError("input-inaccessible") from error

    candidates: list[Path] = []
    for path in entries:
        if not _completed_name(path.name):
            continue
        try:
            entry_stat = path.lstat()
        except FileNotFoundError:
            # The name was present at discovery but disappeared before classification.
            candidates.append(path)
            continue
        except OSError:
            candidates.append(path)
            continue
        if stat.S_ISREG(entry_stat.st_mode):
            candidates.append(path)
    return sorted(candidates, key=lambda item: item.name)


def _same_file_state(before: os.stat_result, after: os.stat_result) -> bool:
    return (
        before.st_dev == after.st_dev
        and before.st_ino == after.st_ino
        and before.st_size == after.st_size
        and before.st_mtime_ns == after.st_mtime_ns
    )


def _read_candidate(path: Path) -> tuple[dict[str, Any] | None, str | None]:
    """Read a candidate once with a strict byte bound and reject detected mutation."""
    try:
        before_path = path.lstat()
        if not stat.S_ISREG(before_path.st_mode):
            return None, "unreadable"
        if before_path.st_size > MAX_INPUT_BYTES:
            return None, "oversized"
        no_follow = getattr(os, "O_NOFOLLOW", 0)
        descriptor = os.open(path, os.O_RDONLY | no_follow)
        with os.fdopen(descriptor, "rb") as stream:
            opened = os.fstat(stream.fileno())
            if not stat.S_ISREG(opened.st_mode) or not _same_file_state(before_path, opened):
                return None, "unreadable"
            if opened.st_size > MAX_INPUT_BYTES:
                return None, "oversized"
            raw = stream.read(MAX_INPUT_BYTES + 1)
            after_open = os.fstat(stream.fileno())
            after_path = path.lstat()
            if not _same_file_state(opened, after_open) or not _same_file_state(opened, after_path):
                return None, "unreadable"
    except OSError:
        return None, "unreadable"
    if len(raw) > MAX_INPUT_BYTES:
        return None, "oversized"
    try:
        parsed = json.loads(raw.decode("utf-8"), object_pairs_hook=_unique_json_object,
                            parse_constant=_reject_nonfinite, parse_int=_parse_json_integer)
    except (UnicodeDecodeError, json.JSONDecodeError, _MalformedMetricJson, ValueError):
        return None, "malformed-json"
    except RecursionError:
        return None, "invalid-contract"
    if (isinstance(parsed, dict) and "schema_version" in parsed
            and (type(parsed["schema_version"]) is not int or parsed["schema_version"] != 1)):
        return None, "unsupported-schema"
    try:
        _require_source_layout()
        record = validate_record(parsed)
    except ValueError:
        return None, "invalid-contract"
    return record, None


def read_metric_records(directory: Path) -> dict[str, Any]:
    """Return validated records and input counts from one fixed discovery snapshot."""
    _require_source_layout()
    candidates = discover_candidates(Path(directory))
    records: list[dict[str, Any]] = []
    invalid_counts: dict[str, int] = {}
    for candidate in candidates:
        record, category = _read_candidate(candidate)
        if category is not None:
            invalid_counts[category] = invalid_counts.get(category, 0) + 1
        elif record is not None:
            records.append(record)
    return {
        "records": records,
        "candidate_count": len(candidates),
        "invalid_counts": {category: invalid_counts[category] for category in INVALID_CATEGORIES if category in invalid_counts},
    }


def consolidate_records(snapshot: dict[str, Any]) -> dict[str, Any]:
    """Collapse identical invocation projections and exclude every conflicting group."""
    _require_source_layout()
    if not isinstance(snapshot, dict) or set(snapshot) != {"records", "candidate_count", "invalid_counts"}:
        raise ValueError("snapshot-contract-invalid")
    records = snapshot["records"]
    candidate_count = snapshot["candidate_count"]
    invalid_counts = snapshot["invalid_counts"]
    if not isinstance(records, list):
        raise ValueError("snapshot-records-invalid")
    if isinstance(candidate_count, bool) or not isinstance(candidate_count, int) or candidate_count < 0:
        raise ValueError("snapshot-candidate-count-invalid")
    if not isinstance(invalid_counts, dict) or any(
        category not in INVALID_CATEGORIES
        or isinstance(count, bool)
        or not isinstance(count, int)
        or count < 0
        for category, count in invalid_counts.items()
    ):
        raise ValueError("snapshot-invalid-counts-invalid")
    normalized_records = [validate_record(record) for record in records]
    invalid_total = sum(invalid_counts.values())
    if candidate_count != len(normalized_records) + invalid_total:
        raise ValueError("snapshot-counts-inconsistent")

    groups: dict[str, list[dict[str, Any]]] = {}
    for record in normalized_records:
        groups.setdefault(record["invocation_id"], []).append(record)

    accepted: list[dict[str, Any]] = []
    duplicate_count = 0
    conflicting_invocation_count = 0
    conflicting_file_count = 0
    for group in groups.values():
        projections = {tuple(record[field] for field in PROJECTION_FIELDS) for record in group}
        if len(projections) != 1:
            conflicting_invocation_count += 1
            conflicting_file_count += len(group)
            continue
        accepted.append(group[0])
        duplicate_count += len(group) - 1

    skipped_count = invalid_total + conflicting_file_count
    summary = {
        "candidate_count": candidate_count,
        "accepted_count": len(accepted),
        "duplicate_count": duplicate_count,
        "conflicting_invocation_count": conflicting_invocation_count,
        "conflicting_file_count": conflicting_file_count,
        "invalid_counts": {category: invalid_counts[category] for category in INVALID_CATEGORIES if category in invalid_counts},
        "skipped_count": skipped_count,
    }
    if summary["accepted_count"] + duplicate_count + skipped_count != candidate_count:
        raise ValueError("snapshot-counts-inconsistent")
    return {"records": accepted, "summary": summary}


def write_workbook(records: list[dict[str, Any]], output_path: Path) -> Path:
    """Write consolidated metric records to one caller-owned XLSX artifact.

    This function creates the workbook at ``output_path``. The caller owns
    temporary naming and publication of that artifact.
    """
    if not isinstance(records, list):
        raise ValueError("workbook-records-invalid")
    _require_source_layout()
    _require_report_dependencies()
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.worksheet.table import Table, TableColumn
    ordered_records = sorted(
        records,
        key=lambda record: (record["started_at"], record["invocation_id"]),
    )
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Metrics")

    header_cells = []
    for header in WORKBOOK_HEADERS:
        cell = WriteOnlyCell(sheet, value=header)
        cell.data_type = "s"
        header_cells.append(cell)
    sheet.append(header_cells)

    for record in ordered_records:
        cells = []
        for field in WORKBOOK_FIELDS:
            value = record[field]
            if field == "duration_ms":
                cell = WriteOnlyCell(sheet, value=value)
            elif value is None:
                cell = WriteOnlyCell(sheet, value=None)
            else:
                cell = WriteOnlyCell(sheet, value=value)
                # openpyxl recognizes formula-looking values during assignment;
                # force the contract's explicit literal-text representation.
                cell.data_type = "s"
            cells.append(cell)
        sheet.append(cells)

    table = Table(displayName="FarmBoyMetrics", ref=f"A1:M{len(ordered_records) + 1}")
    table.tableColumns = [TableColumn(id=index, name=header) for index, header in enumerate(WORKBOOK_HEADERS, 1)]
    sheet.add_table(table)
    destination = Path(output_path)
    workbook.save(destination)
    # XML parsers normalize raw carriage returns in text nodes to line feeds.
    # Encode them as character references in the generated worksheet so the
    # workbook round-trips accepted metric text without changing its value.
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent, delete=False
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        with zipfile.ZipFile(destination, "r") as source, zipfile.ZipFile(temporary_path, "w") as rewritten:
            rewritten.comment = source.comment
            found_worksheet = False
            for entry in source.infolist():
                is_worksheet = entry.filename == "xl/worksheets/sheet1.xml"
                found_worksheet = found_worksheet or is_worksheet
                with source.open(entry, "r") as original_entry, rewritten.open(entry, "w") as copied_entry:
                    while chunk := original_entry.read(64 * 1024):
                        if is_worksheet:
                            # openpyxl emits compact worksheet XML; raw CR bytes
                            # occur only in cell text and must be character references.
                            chunk = chunk.replace(b"\r", b"&#13;")
                        copied_entry.write(chunk)
            if not found_worksheet:
                raise RuntimeError("generated workbook worksheet is missing")
        os.replace(temporary_path, destination)
        temporary_path = None
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
    return destination


def export_workbook(directory: Path | None = None) -> dict[str, Any]:
    """Export one metrics snapshot to a uniquely named, no-clobber workbook."""
    _require_source_layout()
    _require_report_dependencies()
    input_directory = Path(directory) if directory is not None else Path.home() / ".farmboy" / "metrics"
    consolidated = consolidate_records(read_metric_records(input_directory))
    records = consolidated["records"]
    summary = consolidated["summary"]
    if summary["accepted_count"] > MAX_ACCEPTED_ROWS:
        raise ValueError("capacity-exceeded")

    reports_directory = input_directory / "reports"
    try:
        reports_directory.mkdir(parents=True, exist_ok=True)
    except OSError as error:
        raise OSError("report-output-inaccessible") from error

    report_id = uuid.uuid4()
    final_path = reports_directory / f"farmboy-metrics-{report_id}.xlsx"
    stage_path: Path | None = None
    published = False
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".farmboy-metrics-", suffix=".tmp", dir=reports_directory, delete=False
        ) as stage:
            stage_path = Path(stage.name)
        write_workbook(records, stage_path)
        # Hard-link publication is atomic and fails if the UUID destination exists.
        os.link(stage_path, final_path)
        published = True
        stage_path.unlink()
        stage_path = None
    except OSError as error:
        if published and stage_path is not None:
            try:
                if os.path.samefile(stage_path, final_path):
                    final_path.unlink(missing_ok=True)
            except OSError:
                pass
        raise OSError("report-output-failed") from error
    finally:
        if stage_path is not None:
            try:
                stage_path.unlink(missing_ok=True)
            except OSError:
                pass

    return {"output_path": str(final_path), **summary}


def main() -> int:
    """Run the default local report and print bounded status without host paths."""
    try:
        result = export_workbook()
    except (OSError, RuntimeError, ValueError) as error:
        reason = str(error)
        if reason not in {"capacity-exceeded", "input-inaccessible", "report-output-inaccessible", "report-output-failed"}:
            if "layout" in reason:
                reason = "report source layout is incomplete"
            elif "dependenc" in reason:
                reason = "approved report dependency unavailable (openpyxl and et-xmlfile)"
            else:
                reason = "report failed"
        print(f"FarmBoy metrics report failed: {reason}", file=sys.stderr)
        return 1
    output_name = Path(result["output_path"]).name
    print(f"Workbook: ~/.farmboy/metrics/reports/{output_name}")
    for key in ("candidate_count", "accepted_count", "duplicate_count", "conflicting_invocation_count",
                "conflicting_file_count", "skipped_count"):
        print(f"{key}: {result[key]}")
    for category, count in result["invalid_counts"].items():
        print(f"invalid_{category}: {count}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
